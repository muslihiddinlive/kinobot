"""Avto-hisobot: XLSX (Excel diagrammalari bilan). Kunlik/haftalik superadminga yuboriladi."""
import asyncio
import io
import logging
import os
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

from aiogram import Bot
from aiogram.types import BufferedInputFile

import database as db
from config import SUPERADMIN_ID
from utils import admin_data as ad
from utils import search_engine as engine

log = logging.getLogger(__name__)
TZ_OFFSET = int(os.getenv("TZ_OFFSET_HOURS", "5"))      # Toshkent UTC+5
HEAD_FILL = PatternFill("solid", fgColor="2563EB")
HEAD_FONT = Font(bold=True, color="FFFFFF")


def _head(ws, headers):
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(horizontal="center")


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w


def build_xlsx(days: int = 7):
    """-> (bytes, fayl nomi)."""
    d = ad.dashboard(days)
    wb = Workbook()

    ws = wb.active
    ws.title = "Xulosa"
    _head(ws, ["Ko'rsatkich", "Qiymat"])
    for k, v in [
        (f"Hisobot davri (kun)", days), ("Jami foydalanuvchilar", d["total"]), ("Bugun yangi", d["new_today"]),
        (f"Yangi ({days} kun)", d["new_period"]), ("DAU (bugun)", d["dau"]), ("WAU (7 kun)", d["wau"]),
        ("MAU (30 kun)", d["mau"]), ("DAU/MAU (%)", d["stickiness"]), ("Pullik (PRO/VIP) hozir", d["paid_now"]),
        ("Pullik ulushi (%)", d["paid_pct"]), (f"Stars ({days} kun)", d["stars_period"]), ("To'lovlar soni", d["pay_count"]),
        ("To'lovchilar", d["payers"]), ("Konversiya: to'lovchi/MAU (%)", d["conversion"]),
        (f"Qidiruvlar ({days} kun)", d["search_period"]), ("Topilmagan (%)", d["miss_pct"]), ("Banlangan", d["banned"]),
    ]:
        ws.append([k, v])
    _widths(ws, [38, 16])
    if d["activity_since"]:
        ws.append([])
        ws.append([f"Faollik (DAU/WAU/MAU) {d['activity_since']} dan beri o'lchanmoqda"])

    w2 = wb.create_sheet("Kunlik")
    _head(w2, ["Sana", "Yangi userlar", "Faol (DAU)", "Qidiruv", "Stars"])
    nu, dau, se, st = ad.series_new_users(days), ad.series_dau(days), ad.series_searches(days), ad.series_stars(days)
    for i in range(days):
        w2.append([nu[i][0], nu[i][1], dau[i][1], se[i][1], st[i][1]])
    _widths(w2, [14, 14, 12, 12, 12])
    n = days + 1
    cats = Reference(w2, min_col=1, min_row=2, max_row=n)
    b = BarChart(); b.title = "Yangi userlar va Stars"; b.height, b.width = 8, 18
    b.add_data(Reference(w2, min_col=2, min_row=1, max_row=n), titles_from_data=True)
    b.add_data(Reference(w2, min_col=5, min_row=1, max_row=n), titles_from_data=True)
    b.set_categories(cats)
    w2.add_chart(b, "G2")
    l = LineChart(); l.title = "Faol userlar (DAU) va qidiruvlar"; l.height, l.width = 8, 18
    l.add_data(Reference(w2, min_col=3, min_row=1, max_row=n), titles_from_data=True)
    l.add_data(Reference(w2, min_col=4, min_row=1, max_row=n), titles_from_data=True)
    l.set_categories(cats)
    w2.add_chart(l, "G20")

    w3 = wb.create_sheet("Top kinolar")
    _head(w3, ["Kod", "Nomi", "Yuklashlar", "Trending"])
    for m in ad.top_movies(20):
        w3.append([m["code"], m["title"], m["downloads"], "ha" if m["is_trending"] else ""])
    _widths(w3, [14, 40, 12, 10])

    w4 = wb.create_sheet("Topilmaganlar")
    _head(w4, ["So'rov", "Necha marta", "Foydalanuvchilar"])
    for q, c, u in engine.missing_queries(max(days, 30), 50):
        w4.append([q, c, u])
    _widths(w4, [36, 14, 18])

    w5 = wb.create_sheet("To'lovlar")
    _head(w5, ["Sana", "User ID", "Stars", "Reja"])
    conn = db.get_conn()
    for r in conn.execute("SELECT paid_at,user_id,stars,plan FROM payments WHERE date(paid_at)>=? ORDER BY id DESC LIMIT 1000",
                          (nu[0][0],)).fetchall():
        w5.append([r["paid_at"], r["user_id"], r["stars"], r["plan"]])
    conn.close()
    _widths(w5, [20, 14, 10, 10])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"kinobot_hisobot_{datetime.utcnow():%Y%m%d}_{days}k.xlsx"


def summary_caption(days: int) -> str:
    d = ad.dashboard(days)
    return (f"📊 <b>Hisobot ({days} kun)</b>\n\n👥 {d['total']} user (+{d['new_period']})\n"
            f"🟢 DAU {d['dau']} | WAU {d['wau']} | MAU {d['mau']}\n"
            f"💰 ⭐{d['stars_period']} ({d['pay_count']} to'lov)\n"
            f"🔎 Topilmagan: {d['miss_pct']}%")


async def send_report(bot: Bot, days: int, chat_id: int = None) -> None:
    data, name = await asyncio.to_thread(build_xlsx, days)
    await bot.send_document(chat_id or SUPERADMIN_ID, BufferedInputFile(data, name),
                            caption=await asyncio.to_thread(summary_caption, days), parse_mode="HTML")


def _local_now() -> datetime:
    return datetime.utcnow() + timedelta(hours=TZ_OFFSET)


async def _tick(bot: Bot) -> None:
    now = _local_now()
    if now.hour < int(db.get_setting("report_hour", "9") or 9):
        return
    today = str(now.date())
    if db.get_setting("report_daily", "0") == "1" and db.get_setting("report_last_daily", "") != today:
        db.set_setting("report_last_daily", today)         # avval belgilanadi: xato bo'lsa ham spam bo'lmasin
        await send_report(bot, 7)
    if now.weekday() == 0 and db.get_setting("report_weekly", "0") == "1" and db.get_setting("report_last_weekly", "") != today:
        db.set_setting("report_last_weekly", today)
        await send_report(bot, 30)


async def report_loop(bot: Bot) -> None:
    if not SUPERADMIN_ID:
        return
    await asyncio.sleep(120)
    while True:
        try:
            await _tick(bot)
        except asyncio.CancelledError:
            raise
        except Exception as e:  # noqa: BLE001
            log.error("report_loop: %s", e)
        await asyncio.sleep(60)
